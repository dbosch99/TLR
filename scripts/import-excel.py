"""Estrazione in sola lettura del catalogo e delle proprietà dell'acqua dal workbook TLR."""
import hashlib, json, sys, warnings
from pathlib import Path
import openpyxl
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
source = Path(sys.argv[1])
target = Path(__file__).resolve().parents[1] / 'data' / 'catalog.js'
w = openpyxl.load_workbook(source, data_only=True)
s = w['Supporto tubi']; families = {}
for row in range(3, s.max_row + 1):
    kind, dn, inner, outer, mass, fluid, u, ambient = [s.cell(row, c).value for c in range(2, 10)]
    if not kind: continue
    assert all(isinstance(x, (int, float)) and x > 0 for x in [dn, inner, outer, u]), (row, kind)
    steel = kind == 'Acciaio' or kind.startswith('ACCIAIO PREISOLATO')
    family = families.setdefault(kind, {'id': f'pipe-{len(families) + 1}', 'name': kind, 'roughnessMm': w['Dati']['C3' if steel else 'C4'].value, 'ambientC': ambient, 'sizes': []})
    family['sizes'].append({'dn': dn, 'innerMm': inner, 'outerMm': outer, 'uWmK': u})
for family in families.values():
    family['sizes'].sort(key=lambda p: p['innerMm'])
    assert len({x['dn'] for x in family['sizes']}) == len(family['sizes'])
def col(name):
    sheet, area = next(w.defined_names[name].destinations)
    return [c.value for row in w[sheet][area] for c in row]
water = [dict(zip(['temperatureC','densityKgM3','cpKJkgK','viscosityM2s'], point)) for point in zip(col('T_C'), col('ρ_acqua'), col('cp_acqua'), col('Visc'))]
data = {'metadata': {'source': source.name, 'sourceSha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'author': 'Davide Boschetti', 'extractedOn': '2026-10-15', 'note': 'Dati estratti dal workbook di riferimento. Cataloghi incorporati, non aggiornati automaticamente dai produttori.'}, 'water': water, 'families': list(families.values())}
target.write_text('// Dati di riferimento: rigenerare con scripts/import-excel.py, senza modificare il workbook.\nexport const CATALOG = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n')
print(f'{len(families)} famiglie, {sum(len(f["sizes"]) for f in families.values())} diametri, {len(water)} temperature.')
